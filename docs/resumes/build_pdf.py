import os
import subprocess
import re
import shutil
import fitz
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

html_template = """<!DOCTYPE html>
<html lang="en">
<head>
<meta charset="UTF-8">
<title>Shaunak Rane - Resume</title>
<style>
  @page {
    size: letter;
    margin: 0.30in 0.40in 0.26in 0.40in;
  }
  * {
    box-sizing: border-box;
    margin: 0;
    padding: 0;
    font-variant-ligatures: none !important;
    font-feature-settings: "liga" 0, "clig" 0;
  }
  body {
    font-family: Arial, Calibri, 'Segoe UI', sans-serif;
    color: #111111;
    background: #ffffff;
    line-height: 1.25;
    font-size: 8.85pt;
    font-variant-ligatures: none;
    -webkit-font-smoothing: antialiased;
  }
  .header {
    text-align: center;
    margin-bottom: 4px;
  }
  .header h1 {
    font-size: 19pt;
    font-weight: 700;
    letter-spacing: 0.5px;
    color: #000000;
    margin-bottom: 1.5px;
    text-transform: uppercase;
  }
  .header .headline {
    font-size: 9.5pt;
    font-weight: 600;
    color: #1a1a1a;
    margin-bottom: 2.5px;
  }
  .header .contact {
    font-size: 8.2pt;
    color: #333333;
    white-space: nowrap;
  }
  .header .contact a {
    color: #111111;
    text-decoration: underline;
  }
  .section {
    margin-bottom: 4.5px;
  }
  .section-title {
    font-size: 9.6pt;
    font-weight: 700;
    text-transform: uppercase;
    letter-spacing: 0.6px;
    border-bottom: 1px solid #111111;
    padding-bottom: 1px;
    margin-bottom: 3px;
    color: #000000;
  }
  .skills-row {
    font-size: 8.75pt;
    margin-bottom: 1.5px;
    line-height: 1.23;
  }
  .skills-row strong {
    font-weight: 700;
    color: #000000;
  }
  .item {
    margin-bottom: 3.5px;
  }
  .item:last-child {
    margin-bottom: 0;
  }
  .item-header {
    display: flex;
    justify-content: space-between;
    align-items: baseline;
    font-size: 8.95pt;
    margin-bottom: 1px;
  }
  .item-title {
    font-weight: 700;
    color: #000000;
  }
  .item-sub {
    font-style: italic;
    color: #333333;
  }
  .item-date {
    font-size: 8.4pt;
    font-weight: 600;
    color: #222222;
    white-space: nowrap;
  }
  .item-links {
    font-size: 8.2pt;
    color: #333333;
    margin-bottom: 1.5px;
    line-height: 1.20;
  }
  .item-links a {
    color: #111111;
    text-decoration: underline;
  }
  ul.bullets {
    margin-left: 15px;
    list-style-type: disc;
  }
  ul.bullets li {
    font-size: 8.7pt;
    line-height: 1.24;
    margin-bottom: 1.2px;
    text-align: justify;
  }
  ul.bullets li:last-child {
    margin-bottom: 0;
  }
</style>
</head>
<body>

  <div class="header">
    <h1>Shaunak Rane</h1>
    <div class="headline">Software Engineering | Python | Backend Systems | AI/ML</div>
    <div class="contact">
      Thane, India | +91 93202 21211 | 
      <a href="mailto:shaunakrane914@gmail.com">shaunakrane914@gmail.com</a> | 
      <a href="https://www.linkedin.com/in/shaunak-rane-3980582ba/" target="_blank">linkedin.com/in/shaunak-rane</a> | 
      <a href="https://github.com/Shaunakrane914" target="_blank">github.com/Shaunakrane914</a> | 
      <a href="https://shaunak-rane.pages.dev" target="_blank">shaunak-rane.pages.dev</a>
    </div>
  </div>

  <div class="section">
    <div class="section-title">Education</div>
    <div class="item">
      <div class="item-header">
        <div>
          <span class="item-title">Universal AI University</span> - <span class="item-sub">B.Tech in Artificial Intelligence &amp; Machine Learning</span>
        </div>
        <div class="item-date">Expected May 2028 | Karjat, Maharashtra</div>
      </div>
      <ul class="bullets">
        <li><strong>CGPA: 8.47 / 10.0</strong> | <strong>Relevant Coursework:</strong> Data Structures &amp; Algorithms, Object-Oriented Programming, Operating Systems, DBMS, Computer Networks, System Design.</li>
      </ul>
    </div>
  </div>

  <div class="section">
    <div class="section-title">Technical Skills</div>
    <div class="skills-row">
      <strong>Languages &amp; Core:</strong> Python, C, SQL, Bash, Git, Linux (CLI, Scripting)
    </div>
    <div class="skills-row">
      <strong>Backend &amp; Infrastructure:</strong> FastAPI, Flask, RESTful APIs, WebSockets, PostgreSQL, Supabase, MySQL, Docker, Linux Internals
    </div>
    <div class="skills-row">
      <strong>Software Engineering:</strong> Data Structures &amp; Algorithms (DSA), OOP, pytest (Unit &amp; Integration Testing), Alembic (Schema Migrations), CI/CD
    </div>
    <div class="skills-row">
      <strong>AI &amp; Intelligent Systems:</strong> Multi-Agent Workflows, RAG Architecture, PyTorch, PyTorch Geometric, Vector Embeddings (pgvector), Time-Series Modeling
    </div>
  </div>

  <div class="section">
    <div class="section-title">Work Experience</div>
    
    <div class="item">
      <div class="item-header">
        <div>
          <span class="item-title">Univitt Technologies</span> - <span class="item-sub">AI &amp; Systems Engineering Intern</span>
        </div>
        <div class="item-date">May 2026 - Aug 2026 | Remote</div>
      </div>
      <ul class="bullets">
        <li>Engineered predictive Condition-Based Monitoring (CBM) pipeline processing 4,000+ daily historian records across 3 industrial plants, reconciling multi-stage thermodynamics.</li>
        <li>Re-implemented historical Degradation Factor Index (DFI) models across 3,879 historian logs with numerical convergence (&lt;1e-10 residual vs baseline); formulated deterministic maintenance-reset logic validated across 17 automated regression test gates.</li>
        <li>Containerized Flask REST inference microservices with Docker for edge deployment, delivering sub-second anomaly detection and thermodynamic scoring endpoints.</li>
        <li>Audited 400+ physical and mechanical process variables, evaluating ML models against Holt damped-trend statistical baselines; instituted automated CV promotion gates that halted ungrounded maintenance alerts.</li>
      </ul>
    </div>

    <div class="item">
      <div class="item-header">
        <div>
          <span class="item-title">Univitt Technologies</span> - <span class="item-sub">Software Engineering Intern</span>
        </div>
        <div class="item-date">Apr 2025 - Jul 2025 | Remote</div>
      </div>
      <ul class="bullets">
        <li>Developed RESTful backend APIs and normalized relational PostgreSQL database schemas with FastAPI for an institutional food operations and cost-management platform (serving 1,200+ daily meals).</li>
        <li>Automated weekly menu planning and multi-tier Bill-of-Materials (BOM) ingredient rollups, eliminating manual inventory calculation overhead.</li>
        <li>Trained a Random Forest meal-attendance forecaster incorporating weather, calendar events, and shift historical trends, achieving 88% accuracy (within &plusmn;10% margin) and outperforming rolling baselines by 14%.</li>
        <li>Maintained 85%+ backend test coverage with pytest, executed Alembic database schema migrations, and collaborated on code reviews using Git pull-request workflows.</li>
      </ul>
    </div>
  </div>

  <div class="section">
    <div class="section-title">Projects</div>

    <div class="item">
      <div class="item-header">
        <div>
          <span class="item-title">Project Aegis - Threat Intelligence &amp; Claim Verification Engine</span>
        </div>
        <div class="item-date">Python, FastAPI, Multi-Agent Swarms, RAG, pgvector, Docker</div>
      </div>
      <div class="item-links">
        <strong>Live Demo:</strong> <a href="https://aegis-protocol-110.pages.dev" target="_blank">aegis-protocol-110.pages.dev</a> | <strong>Source Code:</strong> <a href="https://github.com/Shaunakrane914/Misinformation" target="_blank">github.com/Shaunakrane914/Misinformation</a>
      </div>
      <ul class="bullets">
        <li>Architected decoupled multi-agent verification pipeline (Scout &rarr; Research &rarr; Investigator) enforcing strict Pydantic JSON Schema validation, achieving zero unhandled schema errors across benchmark test suites via automated retry policies.</li>
        <li>Engineered quantitative market anomaly detection (Scout Agent) computing 5-day rolling Z-scores on price/volume telemetry via Yahoo Finance proxies to isolate abnormal spikes.</li>
        <li>Implemented canonical SHA-256 claim hashing for sub-millisecond in-memory cache deduplication, preventing redundant LLM token spend, paired with Supabase pgvector semantic retrieval and auditable verdict trails.</li>
      </ul>
    </div>

    <div class="item">
      <div class="item-header">
        <div>
          <span class="item-title">TopoFlow - GNN Permeability &amp; Micro-CT Pore Network Benchmark</span>
        </div>
        <div class="item-date">Python, PyTorch Geometric, GraphSAGE, OpenPNM, FastAPI</div>
      </div>
      <div class="item-links">
        <strong>Source Code:</strong> <a href="https://github.com/Shaunakrane914/Flow" target="_blank">github.com/Shaunakrane914/Flow</a>
      </div>
      <ul class="bullets">
        <li>Benchmarked GraphSAGE against classical Kozeny-Carman physics baselines across 1,231 micro-CT pore networks from 5 geological formations, establishing an empirical regime threshold (Cv) for when graph topology outperforms bulk equations.</li>
        <li>Achieved 46.2% MSE reduction on heterogeneous Savonnières carbonate and 28.4% on Estaillades formations over classical solvers, while identifying that homogeneous sandstones remain better predicted by classical physics.</li>
        <li>Built end-to-end pore-network extraction pipeline (PoreSpy/SNOW2), GNN model training workflows in PyTorch Geometric, and streaming FastAPI inference endpoints with SSE.</li>
      </ul>
    </div>

    <div class="item">
      <div class="item-header">
        <div>
          <span class="item-title">Gridium Protocol - Autonomous Microgrid Optimization &amp; Telemetry</span>
        </div>
        <div class="item-date">Python, PyTorch, DDPG, FastAPI, WebSockets, Solidity, Circom</div>
      </div>
      <div class="item-links">
        <strong>Live Demo:</strong> <a href="https://live-ai-1-7tcy.vercel.app" target="_blank">live-ai-1-7tcy.vercel.app</a> | <strong>Source Code:</strong> <a href="https://github.com/Shaunakrane914/Live-Ai-1" target="_blank">github.com/Shaunakrane914/Live-Ai-1</a>
      </div>
      <ul class="bullets">
        <li>Developed continuous-control reinforcement learning engine (PyTorch DDPG) in a custom 15-node Ohm's law Gymnasium microgrid (AegisEnv), dynamically adjusting AMM swap fees (0.10%&ndash;5.00%) to mitigate solar duck-curve volatility.</li>
        <li>Reduced peak-hour grid strain by 18.4% compared to static flat-fee and rule-based fixed-tariff baselines across simulated multi-node solar generation and battery demand cycles.</li>
        <li>Engineered 500ms Socket.io telemetry gateway, Solidity constant-product AMM contract (x &times; y = k) with reentrancy protection, and Groth16 zk-SNARK circuits proving energy surplus off-chain.</li>
      </ul>
    </div>
  </div>

</body>
</html>
"""

def generate_docx(docx_path):
    doc = Document()
    for section in doc.sections:
        section.top_margin = Inches(0.33)
        section.bottom_margin = Inches(0.30)
        section.left_margin = Inches(0.42)
        section.right_margin = Inches(0.42)

    def add_p(text="", space_before=0, space_after=2, line_spacing=1.1, align=WD_ALIGN_PARAGRAPH.LEFT):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(space_before)
        p.paragraph_format.space_after = Pt(space_after)
        p.paragraph_format.line_spacing = line_spacing
        p.alignment = align
        if text:
            p.add_run(text)
        return p

    def add_heading(title):
        p = add_p(space_before=4.5, space_after=2)
        run = p.add_run(title.upper())
        run.bold = True
        run.font.name = "Arial"
        run.font.size = Pt(9.8)
        run.font.color.rgb = RGBColor(0, 0, 0)
        
        pPr = p._p.get_or_add_pPr()
        pBdr = OxmlElement('w:pBdr')
        bottom = OxmlElement('w:bottom')
        bottom.set(qn('w:val'), 'single')
        bottom.set(qn('w:sz'), '6')
        bottom.set(qn('w:space'), '1')
        bottom.set(qn('w:color'), '000000')
        pBdr.append(bottom)
        pPr.append(pBdr)

    # Header
    p_name = add_p(space_before=0, space_after=1, align=WD_ALIGN_PARAGRAPH.CENTER)
    r_name = p_name.add_run("SHAUNAK RANE")
    r_name.bold = True
    r_name.font.name = "Arial"
    r_name.font.size = Pt(18)

    p_head = add_p(space_before=0, space_after=1, align=WD_ALIGN_PARAGRAPH.CENTER)
    r_head = p_head.add_run("Software Engineering | Python | Backend Systems | AI/ML")
    r_head.bold = True
    r_head.font.name = "Arial"
    r_head.font.size = Pt(9.5)

    p_contact = add_p(space_before=0, space_after=3.5, align=WD_ALIGN_PARAGRAPH.CENTER)
    r_con = p_contact.add_run("Thane, India | +91 93202 21211 | shaunakrane914@gmail.com | linkedin.com/in/shaunak-rane | github.com/Shaunakrane914 | shaunak-rane.pages.dev")
    r_con.font.name = "Arial"
    r_con.font.size = Pt(8.5)

    # Education (Near Top)
    add_heading("Education")
    p_edu = add_p(space_before=2, space_after=1)
    r_u = p_edu.add_run("Universal AI University")
    r_u.bold = True
    r_u.font.name = "Arial"
    r_u.font.size = Pt(9.1)
    r_deg = p_edu.add_run(" - B.Tech in Artificial Intelligence & Machine Learning")
    r_deg.font.name = "Arial"
    r_deg.font.size = Pt(9.1)
    r_edate = p_edu.add_run("       (Expected May 2028 | Karjat, Maharashtra)")
    r_edate.font.name = "Arial"
    r_edate.font.size = Pt(8.4)
    r_edate.bold = True

    bp_c = doc.add_paragraph(style='List Bullet')
    bp_c.paragraph_format.space_before = Pt(0)
    bp_c.paragraph_format.space_after = Pt(1)
    br_c = bp_c.add_run("CGPA: ")
    br_c.bold = True
    br_c.font.name = "Arial"
    br_c.font.size = Pt(8.55)
    br_cgpa = bp_c.add_run("8.47 / 10.0")
    br_cgpa.bold = True
    br_cgpa.font.name = "Arial"
    br_cgpa.font.size = Pt(8.55)
    br_sep = bp_c.add_run("   |   Relevant Coursework: ")
    br_sep.font.name = "Arial"
    br_sep.font.size = Pt(8.55)
    br_ci = bp_c.add_run("Data Structures & Algorithms, Object-Oriented Programming, Operating Systems, DBMS, Computer Networks, System Design.")
    br_ci.font.name = "Arial"
    br_ci.font.size = Pt(8.55)

    # Technical Skills
    add_heading("Technical Skills")
    skills = [
        ("Languages & Core: ", "Python, C, SQL, Bash, Git, Linux (CLI, Scripting)"),
        ("Backend & Infrastructure: ", "FastAPI, Flask, RESTful APIs, WebSockets, PostgreSQL, Supabase, MySQL, Docker, Linux Internals"),
        ("Software Engineering: ", "Data Structures & Algorithms (DSA), OOP, pytest (Unit & Integration Testing), Alembic (Schema Migrations), CI/CD"),
        ("AI & Intelligent Systems: ", "Multi-Agent Workflows, RAG Architecture, PyTorch, PyTorch Geometric, Vector Embeddings (pgvector), Time-Series Modeling")
    ]
    for title, items in skills:
        p_sk = add_p(space_before=0, space_after=1)
        r_t = p_sk.add_run(title)
        r_t.bold = True
        r_t.font.name = "Arial"
        r_t.font.size = Pt(8.6)
        r_i = p_sk.add_run(items)
        r_i.font.name = "Arial"
        r_i.font.size = Pt(8.6)

    # Experience
    add_heading("Work Experience")

    def add_exp(company, role, dates, bullets):
        p = add_p(space_before=2, space_after=1)
        r1 = p.add_run(company)
        r1.bold = True
        r1.font.name = "Arial"
        r1.font.size = Pt(9.1)
        r2 = p.add_run(f" - {role}")
        r2.italic = True
        r2.font.name = "Arial"
        r2.font.size = Pt(9.1)
        
        p_date = p.add_run(f"       ({dates})")
        p_date.font.name = "Arial"
        p_date.font.size = Pt(8.4)
        p_date.bold = True

        for b in bullets:
            bp = doc.add_paragraph(style='List Bullet')
            bp.paragraph_format.space_before = Pt(0)
            bp.paragraph_format.space_after = Pt(1)
            bp.paragraph_format.line_spacing = 1.15
            br = bp.add_run(b)
            br.font.name = "Arial"
            br.font.size = Pt(8.55)

    add_exp(
        "Univitt Technologies",
        "AI & Systems Engineering Intern",
        "May 2026 - Aug 2026 | Remote",
        [
            "Engineered predictive Condition-Based Monitoring (CBM) pipeline processing 4,000+ daily historian records across 3 industrial plants, reconciling multi-stage thermodynamics.",
            "Re-implemented historical Degradation Factor Index (DFI) models across 3,879 historian logs with numerical convergence (<1e-10 residual vs baseline); formulated deterministic maintenance-reset logic validated across 17 automated regression test gates.",
            "Containerized Flask REST inference microservices with Docker for edge deployment, delivering sub-second anomaly detection and thermodynamic scoring endpoints.",
            "Audited 400+ physical and mechanical process variables, evaluating ML models against Holt damped-trend statistical baselines; instituted automated CV promotion gates that halted ungrounded maintenance alerts."
        ]
    )

    add_exp(
        "Univitt Technologies",
        "Software Engineering Intern",
        "Apr 2025 - Jul 2025 | Remote",
        [
            "Developed RESTful backend APIs and normalized relational PostgreSQL database schemas with FastAPI for an institutional food operations and cost-management platform (serving 1,200+ daily meals).",
            "Automated weekly menu planning and multi-tier Bill-of-Materials (BOM) ingredient rollups, eliminating manual inventory calculation overhead.",
            "Trained a Random Forest meal-attendance forecaster incorporating weather, calendar events, and shift historical trends, achieving 88% accuracy (within ±10% margin) and outperforming rolling baselines by 14%.",
            "Maintained 85%+ backend test coverage with pytest, executed Alembic database schema migrations, and collaborated on code reviews using Git pull-request workflows."
        ]
    )

    # Projects
    add_heading("Projects")

    def add_proj(name, tech, links_text, bullets):
        p = add_p(space_before=2, space_after=0.5)
        r1 = p.add_run(name)
        r1.bold = True
        r1.font.name = "Arial"
        r1.font.size = Pt(9.1)
        r2 = p.add_run(f" | {tech}")
        r2.font.name = "Arial"
        r2.font.size = Pt(8.4)
        r2.italic = True

        if links_text:
            p_lnk = add_p(space_before=0, space_after=1)
            for part in links_text.split(" | "):
                pieces = part.split(": ")
                if len(pieces) == 2:
                    lbl, url = pieces
                    r_lbl = p_lnk.add_run(f"{lbl}: ")
                    r_lbl.bold = True
                    r_lbl.font.name = "Arial"
                    r_lbl.font.size = Pt(8.2)
                    r_u = p_lnk.add_run(f"{url}   ")
                    r_u.font.name = "Arial"
                    r_u.font.size = Pt(8.2)
                    r_u.underline = True

        for b in bullets:
            bp = doc.add_paragraph(style='List Bullet')
            bp.paragraph_format.space_before = Pt(0)
            bp.paragraph_format.space_after = Pt(1)
            bp.paragraph_format.line_spacing = 1.15
            br = bp.add_run(b)
            br.font.name = "Arial"
            br.font.size = Pt(8.55)

    add_proj(
        "Project Aegis - Threat Intelligence & Claim Verification Engine",
        "Python, FastAPI, Multi-Agent Swarms, RAG, pgvector, Docker",
        "Live Demo: https://aegis-protocol-110.pages.dev | Source Code: https://github.com/Shaunakrane914/Misinformation",
        [
            "Architected decoupled multi-agent verification pipeline (Scout -> Research -> Investigator) enforcing strict Pydantic JSON Schema validation, achieving zero unhandled schema errors across benchmark test suites via automated retry policies.",
            "Engineered quantitative market anomaly detection (Scout Agent) computing 5-day rolling Z-scores on price/volume telemetry via Yahoo Finance proxies to isolate abnormal spikes.",
            "Implemented canonical SHA-256 claim hashing for sub-millisecond in-memory cache deduplication, preventing redundant LLM token spend, paired with Supabase pgvector semantic retrieval and auditable verdict trails."
        ]
    )

    add_proj(
        "TopoFlow - GNN Permeability & Micro-CT Pore Network Benchmark",
        "Python, PyTorch Geometric, GraphSAGE, OpenPNM, FastAPI",
        "Source Code: https://github.com/Shaunakrane914/Flow",
        [
            "Benchmarked GraphSAGE against classical Kozeny-Carman physics baselines across 1,231 micro-CT pore networks from 5 geological formations, establishing an empirical regime threshold (Cv) for when graph topology outperforms bulk equations.",
            "Achieved 46.2% MSE reduction on heterogeneous Savonnières carbonate and 28.4% on Estaillades formations over classical solvers, while identifying that homogeneous sandstones remain better predicted by classical physics.",
            "Built end-to-end pore-network extraction pipeline (PoreSpy/SNOW2), GNN model training workflows in PyTorch Geometric, and streaming FastAPI inference endpoints with SSE."
        ]
    )

    add_proj(
        "Gridium Protocol - Autonomous Microgrid Optimization & Telemetry",
        "Python, PyTorch, DDPG, FastAPI, WebSockets, Solidity, Circom",
        "Live Demo: https://live-ai-1-7tcy.vercel.app | Source Code: https://github.com/Shaunakrane914/Live-Ai-1",
        [
            "Developed continuous-control reinforcement learning engine (PyTorch DDPG) in a custom 15-node Ohm's law Gymnasium microgrid (AegisEnv), dynamically adjusting AMM swap fees (0.10%-5.00%) to mitigate solar duck-curve volatility.",
            "Reduced peak-hour grid strain by 18.4% compared to static flat-fee and rule-based fixed-tariff baselines across simulated multi-node solar generation and battery demand cycles.",
            "Engineered 500ms Socket.io telemetry gateway, Solidity constant-product AMM contract (x * y = k) with reentrancy protection, and Groth16 zk-SNARK circuits proving energy surplus off-chain."
        ]
    )

    doc.save(docx_path)
    print(f"Saved DOCX to {docx_path}")

base_dir = r"c:\Users\Shaunak Rane\Desktop\Projects\Portfolio\docs\resumes"
root_dir = r"c:\Users\Shaunak Rane\Desktop\Projects\Portfolio"

html_path = os.path.join(base_dir, "Shaunak_Rane_Resume.html")
pdf_path = os.path.join(base_dir, "Shaunak_Rane_Resume.pdf")
docx_path = os.path.join(base_dir, "Shaunak_Rane_Resume.docx")

with open(html_path, "w", encoding="utf-8") as f:
    f.write(html_template)

chrome_path = r"C:\Program Files\Google\Chrome\Application\chrome.exe"
cmd = [
    chrome_path,
    "--headless=new",
    "--disable-gpu",
    "--no-pdf-header-footer",
    f"--print-to-pdf={pdf_path}",
    html_path
]
subprocess.run(cmd, check=True)

# Measure PDF
doc = fitz.open(pdf_path)
page_count = len(doc)
page = doc[0]
blocks = page.get_text("blocks")
max_y = max(b[3] for b in blocks if b[4].strip())
bottom_gap = page.rect.height - max_y
print(f"RESULTS: Page Count: {page_count}, Page Height: {page.rect.height}, Last text Y: {max_y:.2f}, Bottom Gap: {bottom_gap:.2f} pt ({bottom_gap/page.rect.height*100:.1f}%)")

# Generate DOCX
generate_docx(docx_path)

# Copy to root and Logikwerk copies
for target in [
    os.path.join(root_dir, "Shaunak_Rane_Resume.pdf"),
    os.path.join(base_dir, "Shaunak_Rane_Logikwerk_AI_Engineer.pdf")
]:
    shutil.copy2(pdf_path, target)
    print(f"Copied PDF to {target}")

for target in [
    os.path.join(root_dir, "Shaunak_Rane_Resume.docx"),
    os.path.join(base_dir, "Shaunak_Rane_Logikwerk_AI_Engineer.docx")
]:
    shutil.copy2(docx_path, target)
    print(f"Copied DOCX to {target}")
